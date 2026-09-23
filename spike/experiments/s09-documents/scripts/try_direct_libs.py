"""Baseline: direct format-specific libraries, no markitdown. Used only to
score fact-recovery for the FINDINGS.md comparison table; not the chosen
extractor (see extract.py)."""
import sys

from pptx import Presentation
import openpyxl
from docx import Document
from pypdf import PdfReader


def extract_pptx(path: str) -> str:
    prs = Presentation(path)
    out = []
    for i, slide in enumerate(prs.slides, 1):
        out.append(f"## Slide {i}")
        for shape in slide.shapes:
            if shape.has_text_frame:
                out.append(shape.text_frame.text)
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text
            if notes.strip():
                out.append(f"Notes: {notes}")
    return "\n".join(out)


def extract_xlsx(path: str) -> str:
    wb = openpyxl.load_workbook(path, data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"## {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                out.append(" | ".join(cells))
    return "\n".join(out)


def extract_docx(path: str) -> str:
    doc = Document(path)
    out = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            out.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(out)


def extract_pdf(path: str) -> str:
    reader = PdfReader(path)
    return "\n".join(page.extract_text() or "" for page in reader.pages)


EXTRACTORS = {
    ".pptx": extract_pptx,
    ".xlsx": extract_xlsx,
    ".docx": extract_docx,
    ".pdf": extract_pdf,
}

if __name__ == "__main__":
    import os

    for path in sys.argv[1:]:
        ext = os.path.splitext(path)[1].lower()
        print(f"=== {path} ===")
        fn = EXTRACTORS.get(ext)
        if not fn:
            print(f"ERROR: no direct extractor for {ext}")
            continue
        try:
            print(fn(path))
        except Exception as e:
            print(f"ERROR: {type(e).__name__}: {e}")
        print()
